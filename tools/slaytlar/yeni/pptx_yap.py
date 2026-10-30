# -*- coding: utf-8 -*-
"""Offline PowerPoint versions of the live lessons (indir/M1.pptx …).

Each slide picture fills a 16:9 slide; the speaker notes keep the lecture
text, drop the live-session (CANLI) instructions and add the right answer
and explanation of the question on that slide. Leaderboard and word-cloud
slides exist only in the live session and are left out.
"""
import json, os, re, sys
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
CANLI = os.path.join(REPO, "canli")
OUT = os.path.join(REPO, "indir")
W, H = Emu(12192000), Emu(6858000)   # 13.333 x 7.5 in


def answer(q):
    t = q["type"]
    if t in ("MCQ", "IMGMCQ", "POLL") and q.get("a") is not None and q.get("opts"):
        letter, opt = "ABCDE"[q["a"]], q["opts"][q["a"]]
        s = "Doğru cevap: %s." % letter if opt.strip() == letter else "Doğru cevap: %s) %s." % (letter, opt)
    elif t == "MATCH":
        s = "Doğru eşleşme: " + "; ".join("%s → %s" % (l, k) for l, k in zip(q["left"], q["key"])) + "."
    elif t == "SORTGROUP":
        s = "Doğru gruplar: " + "; ".join("%s → %s" % (it, q["groups"][g]) for it, g in zip(q["items"], q["key"])) + "."
    elif t == "SCT" and q.get("scale"):
        s = "Uzman cevabı: %s." % q["scale"][q["a"]]
    elif t == "FINDIMG":
        s = "Doğru alan: %s." % q["answer"]
    elif t == "OPEN" and q.get("model"):
        s = "Beklenen cevap: %s" % q["model"]
    else:
        return ""
    if q.get("exp"):
        s += " " + q["exp"]
    return s


def build(lid):
    L = json.load(open(os.path.join(CANLI, "dersler", lid + ".json"), encoding="utf-8"))
    qmap = {q["id"]: q for q in L["questions"]}
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    n = 0
    for s in L["slides"]:
        if not s.get("img"):
            continue
        if s.get("step") and s["step"][0] != s["step"][1]:
            continue          # adım adım slaytın yalnız son (tam) hâli
        n += 1
        sl = prs.slides.add_slide(blank)
        sl.shapes.add_picture(os.path.join(CANLI, s["img"].split("?")[0]), 0, 0, W, H)
        notes = re.sub(r"\n\nCANLI \(.*$", "", s.get("notes", ""), flags=re.S).strip()
        notes = re.sub(r"^\[[^\]]*\]\s*", "", notes)            # time stamp
        notes = notes.replace("ZAMAN KALIRSA; yoksa → ile atla.", "Zaman kalırsa.").strip()
        qid = s.get("q") or s.get("qopen")
        if qid and qid in qmap and not (s.get("qopen") and not s.get("q")):
            a = answer(qmap[qid])
            if a:
                notes += "\n\n" + a
        sl.notes_slide.notes_text_frame.text = s["title"] + "\n\n" + notes
    prs.core_properties.title = L["title"]
    prs.core_properties.author = "Uzm. Dr. M. Olgun Ak"
    prs.core_properties.subject = "4. Sınıf Patoloji · " + L["date"]
    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, lid + ".pptx")
    prs.save(path)
    return {"id": lid, "title": L["title"], "date": L["date"], "slides": n, "mb": round(os.path.getsize(path) / 1e6, 1),
            "file": lid + ".pptx", "quiz": L.get("quiz", "").replace("../", "../")}


def page(info):
    """indir/index.html and the site home page (index.html) list the same decks."""
    import html
    tpl = open(os.path.join(HERE, "indir_sablon.html"), encoding="utf-8").read()
    for out, f, r, url, ust in [
            (os.path.join(OUT, "index.html"), "", "../", "olgunak.github.io/lam-basinda/indir", ""),
            (os.path.join(REPO, "index.html"), "indir/", "", "olgunak.github.io/lam-basinda",
             '<nav class="top-acts"><a href="canli/">← Canlı derse katıl</a><a href="quiz/">Quizler →</a></nav>')]:
        cards = []
        for i in info:
            q = i.get("quiz", "").replace("../", r)
            quiz = ('<a class="btn ghost" href="%s">Quiz: kendini sına</a>' % html.escape(q)) if q else ""
            cards.append('<article class="set"><span class="eyebrow">%s</span><h2>%s</h2><span class="s">%d slayt · %s MB · PowerPoint</span>'
                         '<div class="acts"><a class="btn" href="%s" download>İndir (.pptx)</a>%s</div></article>'
                         % (html.escape(i["date"]), html.escape(i["title"]), i["slides"], str(i["mb"]).replace(".", ","), html.escape(f + i["file"]), quiz))
        t = tpl.replace("<!--DERSLER-->", "\n".join(cards)).replace("{{R}}", r).replace("{{URL}}", url).replace("<!--UST-->", ust)
        if not ust:
            t = t.replace("Lam Başında · Dersler</h1>", "Lam Başında · Dersler</h1>")
        else:
            t = t.replace("<title>Lam Başında · Dersler</title>", "<title>Lam Başında</title>").replace("Lam Başında · Dersler</h1>", "Lam Başında</h1>")
        open(out, "w", encoding="utf-8").write(t)


if __name__ == "__main__":
    ids = sys.argv[1:] or [e["id"] for e in json.load(open(os.path.join(CANLI, "dersler", "index.json"), encoding="utf-8"))]
    info = [build(i) for i in ids]
    json.dump(info, open(os.path.join(OUT, "dersler.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    page(info)
    for i in info:
        print(i["file"], i["slides"], "slayt", i["mb"], "MB")
